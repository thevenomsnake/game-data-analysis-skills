from __future__ import annotations

import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference


SCRIPT_DIR = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import workbook_manifest  # noqa: E402


class WorkbookManifestTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "analysis.xlsx"
        book = Workbook()
        report = book.active
        report.title = "分析展示"
        report.append(["分类", "人数"])
        report.append(["A", 10])
        report.append(["B", 20])
        report["D20"] = "secret-cell-value-must-not-enter-manifest"
        chart = BarChart()
        chart.title = "玩家数 × 分类"
        chart.add_data(Reference(report, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        chart.set_categories(Reference(report, min_col=1, min_row=2, max_row=3))
        report.add_chart(chart, "F2")
        hidden = book.create_sheet("计算中间层")
        hidden.sheet_state = "hidden"
        book.save(self.path)
        book.close()

    def tearDown(self) -> None:
        self.temp.cleanup()

    @staticmethod
    def _move_charts_under_drawings(source: Path, target: Path) -> None:
        with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(target, "w") as outgoing:
            for name in incoming.namelist():
                data = incoming.read(name)
                target_name = name.replace("xl/charts/", "xl/drawings/charts/")
                if name == "[Content_Types].xml":
                    data = data.replace(b"/xl/charts/", b"/xl/drawings/charts/")
                elif name.startswith("xl/drawings/_rels/"):
                    data = data.replace(b"/xl/charts/", b"/xl/drawings/charts/")
                    data = data.replace(b"../charts/", b"charts/")
                outgoing.writestr(target_name, data)

    def test_manifest_exposes_only_bounded_presentation_structure(self) -> None:
        manifest = workbook_manifest.build_workbook_manifest(self.path)

        self.assertEqual(manifest["schema_version"], "workbook_manifest_v1")
        self.assertEqual(manifest["sheet_count"], 2)
        self.assertEqual(
            manifest["sheets"],
            [
                {"name": "分析展示", "visibility": "visible"},
                {"name": "计算中间层", "visibility": "hidden"},
            ],
        )
        self.assertEqual(manifest["chart_count"], 1)
        self.assertEqual(manifest["chart_titles"], ["玩家数 × 分类"])
        self.assertNotIn("secret-cell-value", json.dumps(manifest, ensure_ascii=False))
        self.assertTrue(manifest["display_metadata"]["bounded"])

    def test_manifest_recognizes_artifact_tool_chart_location(self) -> None:
        artifact_tool_path = Path(self.temp.name) / "artifact-tool.xlsx"
        self._move_charts_under_drawings(self.path, artifact_tool_path)

        manifest = workbook_manifest.build_workbook_manifest(artifact_tool_path)

        self.assertEqual(manifest["chart_count"], 1)
        self.assertEqual(manifest["chart_titles"], ["玩家数 × 分类"])

    def test_reusable_xlsx_and_result_xlsx_are_distinct_surfaces(self) -> None:
        media = workbook_manifest.XLSX_MEDIA_TYPE
        reusable = workbook_manifest.reusable_workbook_presentation(
            "visualization",
            media,
            "runs/example/visual.xlsx",
            {"workbook_manifest": workbook_manifest.build_workbook_manifest(self.path)},
        )
        result = workbook_manifest.reusable_workbook_presentation(
            "result",
            media,
            "runs/example/result.xlsx",
            {},
        )
        html_visual = workbook_manifest.reusable_workbook_presentation(
            "visualization",
            "text/html",
            "runs/example/visual.html",
            {},
        )

        self.assertTrue(reusable["eligible"])
        self.assertEqual(reusable["preview_status"], "not_available")
        self.assertFalse(result["eligible"])
        self.assertFalse(html_visual["eligible"])

    def test_legacy_reusable_workbook_without_manifest_remains_downloadable(self) -> None:
        presentation = workbook_manifest.reusable_workbook_presentation(
            "analysis_workbook",
            workbook_manifest.XLSX_MEDIA_TYPE,
            "query_workspace/history/legacy-analysis.xlsx",
            {},
        )

        self.assertTrue(presentation["eligible"])
        self.assertEqual(presentation["preview_status"], "not_available")
        self.assertEqual(presentation["workbook_manifest"], {})
        self.assertEqual(
            presentation["download_path"],
            "query_workspace/history/legacy-analysis.xlsx",
        )

    def test_unsafe_workbook_path_is_not_exposed_to_consumers(self) -> None:
        presentation = workbook_manifest.reusable_workbook_presentation(
            "analysis_workbook",
            workbook_manifest.XLSX_MEDIA_TYPE,
            "C:/outside/analysis.xlsx",
            {},
        )

        self.assertFalse(presentation["eligible"])
        self.assertEqual(presentation["download_path"], "")


if __name__ == "__main__":
    unittest.main()
